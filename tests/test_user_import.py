from io import BytesIO

from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.models import User
from app.services.user_import import HEADERS
from tests.conftest import auth


def xlsx(*rows) -> bytes:
    wb = Workbook()
    wb.active.append(HEADERS)
    for row in rows:
        wb.active.append(list(row))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(client, user, content: bytes):
    return client.post("/users/import", files={"file": ("users.xlsx", content)}, headers=auth(user))


def test_import_creates_and_updates(client, db, teacher, student):
    content = xlsx(
        ["new1@bssm.hs.kr", "박마고", "학생", 1, 2, 7],
        [student.email, "김부마", "학생", 3, 1, 3],
        ["t2@bssm.hs.kr", "최선생", "교사", None, None, None, "교무부장", "010-1234-5678"],
    )
    res = upload(client, teacher, content)
    assert res.status_code == 200
    assert res.json() == {"created": 2, "updated": 1}

    db.expire_all()
    assert db.get(User, student.id).grade == 3
    new = db.scalar(select(User).where(User.email == "new1@bssm.hs.kr"))
    assert new.student_id == "1207"


def test_import_rejects_whole_file_on_error(client, db, teacher):
    content = xlsx(
        ["ok@bssm.hs.kr", "정상", "학생", 1, 1, 1],
        ["bad-email", "오류", "학생", 1, 1, 2],
        ["x@bssm.hs.kr", "구분오류", "학부모", None, None, None],
        ["y@bssm.hs.kr", "반없음", "학생", 1, None, 3],
    )
    res = upload(client, teacher, content)
    assert res.status_code == 422
    assert [e.split(":")[0] for e in res.json()["detail"]] == ["3행", "4행", "5행"]
    assert db.scalar(select(User).where(User.email == "ok@bssm.hs.kr")) is None


def test_import_requires_teacher(client, student):
    assert upload(client, student, xlsx()).status_code == 403


def test_import_rejects_non_excel(client, teacher):
    res = upload(client, teacher, b"not excel")
    assert res.status_code == 422


def test_template_download(client, teacher):
    res = client.get("/users/import/template", headers=auth(teacher))
    assert res.status_code == 200
    header = next(load_workbook(BytesIO(res.content)).active.iter_rows(values_only=True))
    assert list(header) == HEADERS
