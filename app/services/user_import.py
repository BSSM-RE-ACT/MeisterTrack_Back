"""엑셀(.xlsx)로 학생·교사 계정을 일괄 등록한다. 이메일이 같으면 정보를 덮어쓴다."""

from dataclasses import dataclass, field
from io import BytesIO

from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, EmailStr, ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Role, User

HEADERS = ["이메일", "이름", "구분", "학년", "반", "번호", "직위", "연락처"]
ROLE_NAMES = {"학생": Role.student, "교사": Role.teacher}


class ImportRow(BaseModel):
    email: EmailStr
    name: str
    role: Role
    grade: int | None = None
    class_no: int | None = None
    number: int | None = None
    position: str | None = None
    phone: str | None = None


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    errors: list[str] = field(default_factory=list)


def build_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "사용자"
    ws.append(HEADERS)
    ws.append(["student1@bssm.hs.kr", "김부마", "학생", 2, 1, 3, None, None])
    ws.append(["teacher@bssm.hs.kr", "이선생", "교사", None, None, None, "담임교사", "010-0000-0000"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse(content: bytes) -> tuple[list[ImportRow], list[str]]:
    try:
        ws = load_workbook(BytesIO(content), read_only=True, data_only=True).active
    except Exception:
        return [], ["엑셀(.xlsx) 파일을 읽을 수 없습니다"]

    rows = ws.iter_rows(values_only=True)
    header = [_cell(v) for v in next(rows, [])]
    missing = [h for h in HEADERS[:3] if h not in header]
    if missing:
        return [], [f"필수 열이 없습니다: {', '.join(missing)}"]
    index = {h: header.index(h) for h in HEADERS if h in header}

    parsed, errors, seen = [], [], set()
    for line, values in enumerate(rows, start=2):
        raw = {h: _cell(values[i]) if i < len(values) else None for h, i in index.items()}
        if not any(raw.values()):
            continue
        role = ROLE_NAMES.get(raw.get("구분") or "")
        if role is None:
            errors.append(f"{line}행: 구분은 '학생' 또는 '교사'여야 합니다")
            continue
        try:
            row = ImportRow(
                email=raw.get("이메일"),
                name=raw.get("이름"),
                role=role,
                grade=raw.get("학년"),
                class_no=raw.get("반"),
                number=raw.get("번호"),
                position=raw.get("직위"),
                phone=raw.get("연락처"),
            )
        except ValidationError as e:
            fields = ", ".join(str(err["loc"][0]) for err in e.errors())
            errors.append(f"{line}행: 값이 올바르지 않습니다 ({fields})")
            continue
        if row.role == Role.student and None in (row.grade, row.class_no, row.number):
            errors.append(f"{line}행: 학생은 학년·반·번호가 필요합니다")
            continue
        if row.email in seen:
            errors.append(f"{line}행: 이메일이 파일 안에서 중복됩니다")
            continue
        seen.add(row.email)
        parsed.append(row)
    return parsed, errors


def import_users(db: Session, content: bytes) -> ImportResult:
    """한 행이라도 오류가 있으면 아무것도 저장하지 않는다."""
    rows, errors = _parse(content)
    result = ImportResult(errors=errors)
    if errors:
        return result

    existing = {u.email: u for u in db.scalars(select(User).where(User.email.in_([r.email for r in rows])))}
    for row in rows:
        user = existing.get(row.email)
        if user is None:
            db.add(User(**row.model_dump()))
            result.created += 1
        else:
            for key, value in row.model_dump().items():
                setattr(user, key, value)
            result.updated += 1
    db.commit()
    return result
